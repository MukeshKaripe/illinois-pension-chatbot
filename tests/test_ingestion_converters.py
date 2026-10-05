"""Each converter's only job: produce text that looks like ILGA's .txt."""

from __future__ import annotations

import io

import pytest
from _pdf import make_pdf
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pension_chat.ingestion.converters import (
    ConversionFailedError,
    NoTextLayerError,
    UnsupportedFormatError,
    to_text,
)
from pension_chat.ingestion.sources import SourceFile


def _file(name: str, data: bytes) -> SourceFile:
    return SourceFile(name, data, "v1")


def _docx(build) -> bytes:  # type: ignore[no-untyped-def]
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _nonblank(text: str) -> list[str]:
    return [line.strip() for line in text.split("\n") if line.strip()]


def test_txt_with_bom_and_crlf() -> None:
    data = "﻿(40 ILCS 5/8-101)\r\nSec. 8-101. Creation of fund.\r\n".encode()
    assert to_text(_file("article_8.txt", data)) == "(40 ILCS 5/8-101)\nSec. 8-101. Creation of fund.\n"


def test_extension_is_case_insensitive() -> None:
    data = b"(40 ILCS 5/8-101)\nSec. 8-101. Creation of fund.\n"
    assert to_text(_file("ARTICLE_8.TXT", data)) == to_text(_file("article_8.txt", data))


def test_md_is_stripped_of_markup() -> None:
    data = b"## (40 ILCS 5/8-101)\n**Sec. 8-101.** Creation of fund.\n"
    assert _nonblank(to_text(_file("a.md", data))) == ["(40 ILCS 5/8-101)", "Sec. 8-101. Creation of fund."]


def test_unknown_extension_is_refused_by_name() -> None:
    with pytest.raises(UnsupportedFormatError, match="scan.tiff"):
        to_text(_file("scan.tiff", b"..."))


def test_docx_text_in_nested_tables_is_extracted_in_order() -> None:
    def build(doc):  # type: ignore[no-untyped-def]
        outer = doc.add_table(rows=1, cols=1)
        cell = outer.cell(0, 0)
        cell.paragraphs[0].text = "\xa0\xa0(40 ILCS 5/8-101)\nSec. 8-101. Creation of fund."
        inner = cell.add_table(rows=1, cols=1)
        inner.cell(0, 0).paragraphs[0].text = "(a) Text held in a table inside a table."
        cell.add_paragraph("(Source: Laws 1963, p. 161.)")

    assert _nonblank(to_text(_file("art8.docx", _docx(build)))) == [
        "(40 ILCS 5/8-101)",
        "Sec. 8-101. Creation of fund.",
        "(a) Text held in a table inside a table.",
        "(Source: Laws 1963, p. 161.)",
    ]


def test_docx_paragraphs_and_tables_keep_document_order() -> None:
    def build(doc):  # type: ignore[no-untyped-def]
        doc.add_paragraph("(40 ILCS 5/Art. 8 heading)")
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].text = "(40 ILCS 5/8-101)"
        doc.add_paragraph("Sec. 8-101. Creation of fund.")

    assert _nonblank(to_text(_file("a.docx", _docx(build)))) == [
        "(40 ILCS 5/Art. 8 heading)",
        "(40 ILCS 5/8-101)",
        "Sec. 8-101. Creation of fund.",
    ]


def test_docx_merged_cell_text_appears_once() -> None:
    def build(doc):  # type: ignore[no-untyped-def]
        table = doc.add_table(rows=1, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.paragraphs[0].text = "(40 ILCS 5/8-101)"

    assert _nonblank(to_text(_file("a.docx", _docx(build)))) == ["(40 ILCS 5/8-101)"]


def test_docx_text_the_walker_cannot_reach_fails_the_file() -> None:
    # A content control at body level holds text the walker does not visit. The
    # completeness check must notice rather than silently drop it — the same
    # failure as nested tables, which cost ~3,300 words of Article 8.
    def build(doc):  # type: ignore[no-untyped-def]
        doc.add_paragraph("(40 ILCS 5/8-101)")
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r>'
            "<w:t>Text inside a content control.</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
        doc.element.body.insert(0, sdt)

    with pytest.raises(ConversionFailedError, match="characters"):
        to_text(_file("a.docx", _docx(build)))


LONG = "and further text so that the page carries a real text layer of some length"


def test_pdf_lines_come_back_joined_and_without_furniture() -> None:
    pdf = make_pdf([
        ["Illinois Compiled Statutes", "(40 ILCS 5/8-101)",
         "Sec. 8-101. Creation of fund. The fund shall be admin-", "Page 1 of 3"],
        ["Illinois Compiled Statutes", "istered by the board " + LONG,
         "(Source: Laws 1963, p. 161.)", "Page 2 of 3"],
        ["Illinois Compiled Statutes", "(40 ILCS 5/8-102)",
         "Sec. 8-102. Terms defined. " + LONG, "Page 3 of 3"],
    ])
    assert _nonblank(to_text(_file("a.pdf", pdf))) == [
        "(40 ILCS 5/8-101)",
        "Sec. 8-101. Creation of fund. The fund shall be administered by the board " + LONG,
        "(Source: Laws 1963, p. 161.)",
        "(40 ILCS 5/8-102)",
        "Sec. 8-102. Terms defined. " + LONG,
    ]


def test_pdf_without_a_text_layer_is_refused() -> None:
    with pytest.raises(NoTextLayerError, match="scan.pdf"):
        to_text(_file("scan.pdf", make_pdf([[], [], ["(40 ILCS 5/8-101)"]])))


def test_a_scanned_page_in_the_middle_refuses_the_pdf() -> None:
    # One scanned page in four is below any majority threshold, and on a first
    # ingestion nothing in the index would notice the section it held going missing.
    pdf = make_pdf([
        ["(40 ILCS 5/8-101)", "Sec. 8-101. Creation of fund. " + LONG],
        ["Sec. 8-101 continued. " + LONG],
        [],
        ["(40 ILCS 5/8-103)", "Sec. 8-103. Fund. " + LONG],
    ])
    with pytest.raises(NoTextLayerError, match="page 3"):
        to_text(_file("partly-scanned.pdf", pdf))


def test_a_blank_cover_and_end_page_are_allowed() -> None:
    pdf = make_pdf([
        [],
        ["(40 ILCS 5/8-101)", "Sec. 8-101. Creation of fund. " + LONG],
        [],
    ])
    assert "(40 ILCS 5/8-101)" in to_text(_file("with-cover.pdf", pdf))


NBSP4 = "\xa0" * 4


def _cells_docx(*paragraphs: str) -> bytes:
    """Paragraphs laid out the way the client's document is: one per table cell,
    indented with non-breaking spaces."""
    def build(doc):  # type: ignore[no-untyped-def]
        for text in paragraphs:
            doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].text = text
    return _docx(build)


def test_docx_items_indented_deeper_than_the_body_are_marked_nested() -> None:
    # The client's document indents nested list items 8 spaces against a 4-space
    # body — where ILGA's text writes "* (a)". Marked the same way, they stay
    # inside their parent instead of becoming top-level subsections.
    data = _cells_docx(
        f"{NBSP4}(40 ILCS 5/8-154)\n{NBSP4}Sec. 8-154. Maximum annuities.",
        f"{NBSP4}(1) The annuities are subject to the following limitations:",
        f"{NBSP4}{NBSP4}(a) No annuity in excess of 60% of the highest salary.",
        f"{NBSP4}{NBSP4}(b) No annuity to a widow in excess of 40%.",
        "(Source: Laws 1963, p. 161.)",
    )
    assert _nonblank(to_text(_file("a.docx", data))) == [
        "(40 ILCS 5/8-154)",
        "Sec. 8-154. Maximum annuities.",
        "(1) The annuities are subject to the following limitations:",
        "* (a) No annuity in excess of 60% of the highest salary.",
        "* (b) No annuity to a widow in excess of 40%.",
        "(Source: Laws 1963, p. 161.)",
    ]


def test_docx_top_level_subsections_are_left_alone() -> None:
    data = _cells_docx(
        f"{NBSP4}(40 ILCS 5/8-115)\n{NBSP4}Sec. 8-115. Future entrant.",
        f"{NBSP4}(a) Any employee who first enters service.",
        f"{NBSP4}(b) Any person who becomes an employee later.",
    )
    assert _nonblank(to_text(_file("a.docx", data)))[2:] == [
        "(a) Any employee who first enters service.",
        "(b) Any person who becomes an employee later.",
    ]


def test_docx_marker_wrapped_onto_an_unindented_line_is_joined_back() -> None:
    # "(d) ... for which he received" / "(i) a disability benefit ..., or (ii) ...":
    # a list inside a sentence, wrapped onto a new, unindented line. Left alone
    # the parser reads "(i)" as a subsection after (d).
    data = _cells_docx(
        f"{NBSP4}(40 ILCS 5/8-226)\n{NBSP4}Sec. 8-226. Computation of service.",
        f"{NBSP4}(d) Any period of disability for which he received\n"
        "(i) a disability benefit under this Article, or (ii) a temporary total disability benefit.",
    )
    assert _nonblank(to_text(_file("a.docx", data)))[2:] == [
        "(d) Any period of disability for which he received (i) a disability benefit under this "
        "Article, or (ii) a temporary total disability benefit.",
    ]
